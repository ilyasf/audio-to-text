"""
Format a transcript via Ollama → .docx

Setup (once):
    pip install python-docx requests

Usage:
    python format-file.py transcript.txt
    python format-file.py transcript.txt --model qwen2.5:32b
    python format-file.py transcript.txt --chunk-minutes 10

Requirements:
    - Ollama running (WSL or Windows), available at localhost:11434
    - Any Russian-capable model (qwen2.5, llama3.1, mistral, etc.)

How to check available models:
    ollama list
"""

import re
import sys
import argparse
import requests
from pathlib import Path


OLLAMA_URL = "http://localhost:11434/api/generate"

SYSTEM_PROMPT = """You are a transcript editor. Your task is to format a fragment of a speech transcript while strictly preserving the original text.

Rules (follow strictly):
1. Do NOT change words, do not paraphrase, do not shorten — formatting only
2. Split the text into paragraphs by meaning (one thought — one paragraph)
3. When a new topic begins, add a heading on its own line in the format "## Topic name"
4. If the text contains a question from someone (another participant, the audience), format it as:
   "Question: <question text>"
   "Answer: <answer text>"
5. Fix only obvious punctuation errors and capitalization at the start of sentences
6. Do NOT remove words, do NOT add anything yourself, do NOT write explanations
7. Return only the formatted text"""


def list_ollama_models() -> list[str]:
    try:
        r = requests.get("http://localhost:11434/api/tags", timeout=5)
        return [m["name"] for m in r.json().get("models", [])]
    except Exception:
        return []


def call_ollama(text: str, model: str) -> str:
    # num_ctx must cover both the input text and the output response
    # ~2.5 characters per token for Russian text
    needed_ctx = max(8192, int(len(text) / 2.5) + 2048)
    payload = {
        "model": model,
        "system": SYSTEM_PROMPT,
        "prompt": text,
        "stream": False,
        "options": {
            "num_ctx": needed_ctx,
            "temperature": 0,   # deterministic output, less "creativity"
        },
    }
    try:
        r = requests.post(OLLAMA_URL, json=payload, timeout=300)
        r.raise_for_status()
        return r.json()["response"].strip()
    except requests.exceptions.ConnectionError:
        print("❌ Ollama is not available on localhost:11434. Make sure it is running.")
        sys.exit(1)


def parse_transcript(txt_path: Path) -> list[dict]:
    """Parse the format [HH:MM:SS --> HH:MM:SS] text"""
    pattern = re.compile(r"\[(\d{2}:\d{2}:\d{2}) --> (\d{2}:\d{2}:\d{2})\] (.+)")
    segments = []
    for line in txt_path.read_text(encoding="utf-8").splitlines():
        m = pattern.match(line.strip())
        if m:
            segments.append({"start": m.group(1), "end": m.group(2), "text": m.group(3)})
        elif line.strip() and segments:
            # line without a timestamp — just text
            segments.append({"start": "", "end": "", "text": line.strip()})
    return segments


def chunk_segments(segments: list[dict], chunk_minutes: int) -> list[tuple[str, list[dict]]]:
    """Split into chunks of N minutes"""
    def to_seconds(ts: str) -> int:
        if not ts:
            return 0
        h, m, s = map(int, ts.split(":"))
        return h * 3600 + m * 60 + s

    chunks = []
    current_chunk = []
    chunk_start = None
    chunk_limit = chunk_minutes * 60

    for seg in segments:
        t = to_seconds(seg["start"])
        if chunk_start is None:
            chunk_start = t
        if t - chunk_start >= chunk_limit and current_chunk:
            label = f"{_fmt(chunk_start)} – {_fmt(t)}"
            chunks.append((label, current_chunk))
            current_chunk = []
            chunk_start = t
        current_chunk.append(seg)

    if current_chunk:
        t = to_seconds(current_chunk[-1]["end"] or current_chunk[-1]["start"])
        label = f"{_fmt(chunk_start)} – {_fmt(t)}"
        chunks.append((label, current_chunk))

    return chunks


def _fmt(seconds: int) -> str:
    m, s = divmod(seconds, 60)
    h, m = divmod(m, 60)
    return f"{h:02d}:{m:02d}:{s:02d}"


def build_docx(structured_chunks: list[tuple[str, str]], output_path: Path):
    from docx import Document
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()

    # Document title
    title = doc.add_heading("Transcript", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for time_label, structured_text in structured_chunks:
        # Chunk timestamp as a subheading
        if time_label:
            p = doc.add_paragraph()
            run = p.add_run(f"⏱ {time_label}")
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x80, 0x80, 0x80)

        for line in structured_text.splitlines():
            line = line.strip()
            if not line:
                continue
            if line.startswith("## "):
                doc.add_heading(line[3:], level=2)
            elif line.startswith("# "):
                doc.add_heading(line[2:], level=1)
            elif line.startswith("Question:"):
                p = doc.add_paragraph()
                run = p.add_run("Question:")
                run.bold = True
                run.font.color.rgb = RGBColor(0x1F, 0x5C, 0x99)
                p.add_run(line[len("Question:"):])
            elif line.startswith("Answer:"):
                p = doc.add_paragraph()
                run = p.add_run("Answer:")
                run.bold = True
                p.add_run(line[len("Answer:"):])
            else:
                doc.add_paragraph(line)

        doc.add_paragraph("")  # spacing between chunks

    doc.save(str(output_path))


def format_transcript(txt_path: str, model: str, chunk_minutes: int, output_path: str = None):
    txt_path = Path(txt_path)
    if not txt_path.exists():
        print(f"File not found: {txt_path}")
        sys.exit(1)

    if output_path is None:
        output_path = txt_path.with_suffix(".docx")
    else:
        output_path = Path(output_path)

    # Check available models
    models = list_ollama_models()
    if models:
        print(f"Available Ollama models: {', '.join(models)}")
        if model not in models:
            print(f"⚠️  Model '{model}' not found. Use one of: {', '.join(models)}")
            sys.exit(1)
    else:
        print("⚠️  Could not retrieve the model list. Continuing with the specified model...")

    segments = parse_transcript(txt_path)
    if not segments:
        print("Could not parse the file. Make sure it is the output of transcribe.py")
        sys.exit(1)

    chunks = chunk_segments(segments, chunk_minutes)
    print(f"Split into {len(chunks)} chunks of {chunk_minutes} min.\n")

    structured_chunks = []
    for i, (label, segs) in enumerate(chunks, 1):
        raw_text = " ".join(s["text"] for s in segs)
        print(f"[{i}/{len(chunks)}] Processing {label} ({len(raw_text)} chars)...")
        structured = call_ollama(raw_text, model)
        structured_chunks.append((label, structured))

    print(f"\nCreating {output_path.name}...")
    build_docx(structured_chunks, output_path)
    print(f"Done! → {output_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Transcript txt → formatted docx via Ollama")
    parser.add_argument("txt", help="Path to the .txt file from transcribe.py")
    parser.add_argument("--model", default="qwen2.5:7b", help="Ollama model (default: qwen2.5:7b)")
    parser.add_argument("--chunk-minutes", type=int, default=15, help="Chunk size in minutes (default: 15)")
    parser.add_argument("--output", help="Path to the output .docx file")
    args = parser.parse_args()

    format_transcript(args.txt, args.model, args.chunk_minutes, args.output)